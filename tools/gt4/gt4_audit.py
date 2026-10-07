"""
Audit the owner's GT4 Specs.xlsx (THE basis of every stock car value) against
what the game carries: RG2's gt4Database.ts GT4_SPECS (what rg2_cars.json was
baked from) and Resources/rg2_cars.json itself.

Excel turned several "a / b" and "a : b" columns into DATES and TIMES. They
decode losslessly here (see cell()):
  Weight Distribution  timedelta -> "hours : minutes"   (2 days, 0:52 = 48 : 52)
  Stiffness, LSD, DF   datetime  -> "month / day"       (2026-02-02 = 2 / 2)
                       year 2000 and day 1 -> "month / 0" ("5 / 0" was read as May 2000)

Usage:  py tools/gt4/gt4_audit.py ["path/to/GT4 Specs.xlsx"]
Writes tools/gt4/gt4_specs.json (every row of every sheet, decoded) and
prints one line per compared field: checked / differ / first examples.
Needs openpyxl (py -m pip install openpyxl). Reads a COPY of the workbook,
so it works while Excel has the file open.

Findings 2026-10-07 (317 game cars):
- Match exactly: names, year, drivetrain, redline, width, front/rear track,
  power, torque-curve points, displacement, engine type.
- RG2's extraction ZEROED the LSD values Excel had turned into dates (64 cars),
  and read the "5'(0 / 0)" cells as 50 / 0. The game does not read LSD yet.
- The game treats the sheet's PS as mechanical hp (engines about 1.4% strong).
- The sheet has only the Stage-3 weight; stock kg = Stage-3 / 0.85 (RG2's estimate).
- Not read by the game at all: weight distribution (every car is 50:50),
  wheelbase (taken from the shared 3D shell), grip F/R, downforce, springs,
  ride height, stiffness, LSD, tyre sizes, inertias, engine brake, yaw radius,
  rev limiter (game = redline + 500), power band.
"""
import os, shutil, sys, tempfile
SRC = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\mcgee\OneDrive\Documents\Game Development\PSX Assets\PSX Racing\GT4 Specs.xlsx"
HERE = os.path.dirname(os.path.abspath(__file__))
_tmp = os.path.join(tempfile.gettempdir(), "gt4_audit_copy.xlsx")
try:
    shutil.copyfile(SRC, _tmp)
except PermissionError:
    # Excel holds the file open with a lock Python's open() cannot share; the shell's copy can.
    import subprocess
    subprocess.run(["powershell", "-NoProfile", "-Command",
                    "Copy-Item -LiteralPath '%s' -Destination '%s' -Force" % (SRC, _tmp)], check=True)
import openpyxl, datetime, re, json, io
CAR_SHEETS=["OLD","- 99","00 - "]
def cell(v):
    if isinstance(v,datetime.timedelta):
        m=int(v.total_seconds()//60); return f"{m//60} : {m%60}"
    if isinstance(v,datetime.datetime):
        if v.year==2000 and v.day==1: return f"{v.month} / 0"
        return f"{v.month} / {v.day}"
    if isinstance(v,float) and v.is_integer(): return int(v)
    return v
def sheet():
    wb=openpyxl.load_workbook(_tmp,data_only=True)
    cars={}; eng={}; sus={}
    for sn in CAR_SHEETS:
        ws=wb[sn]; hdr=[c.value for c in ws[2]]
        for r in ws.iter_rows(min_row=3,values_only=True):
            if not r[1]: continue
            d={h:cell(v) for h,v in zip(hdr,r) if h}; d['_sheet']=sn
            cars[str(r[1]).strip()]=d
    for sn,dst in (("Engine Specs",eng),("Suspension",sus)):
        ws=wb[sn]; hdr=[c.value for c in ws[2]]
        for i,h in enumerate(hdr):
            if h is None and i>0: hdr[i]=f"_c{i}"
        for r in ws.iter_rows(min_row=3,values_only=True):
            if not r[1]: continue
            dst[str(r[1]).strip()]={h:cell(v) for h,v in zip(hdr,r) if h}
    return cars,eng,sus
def rg2():
    src=io.open(r"C:\Users\mcgee\code\Racing-Game-2\src\config\cars\gt4Database.ts",encoding="utf-8").read()
    db={}
    for m in re.finditer(r"^\['((?:[^'\\]|\\.)*)',(\d+),(\d+),'(\w+)',",src,re.M):
        db[m.group(1).replace("\\'","'")]=dict(hp=int(m.group(2)),kg=int(m.group(3)),drv=m.group(4))
    specs={}
    for m in re.finditer(r"^\s*'((?:[^'\\]|\\.)*)':\{([^\n]*)\},?\s*$",src,re.M):
        body=m.group(2); s={}
        tc=re.search(r"\btc:\[(.*?)\],susp:",body)
        if tc: s['tc']=json.loads("["+tc.group(1)+"]"); body=body.replace(tc.group(0),"susp:")
        for k,v in re.findall(r"(\w+):(\[[^\]]*\]|'[^']*'|-?[\d.]+)",body):
            if v.startswith("["): s[k]=json.loads(v)
            elif v.startswith("'"): s[k]=v[1:-1]
            else: s[k]=float(v)
        specs[m.group(1).replace("\\'","'")]=s
    return db,specs
def game():
    return {c['name']:c for c in json.load(open(r"C:\Users\mcgee\PSX Racing\Assets\PSXRacing\Resources\rg2_cars.json",encoding="utf-8"))['cars']}

import collections, re
cars,eng,sus=sheet(); db,specs=rg2(); g=game()
def pair(v):
    if v is None: return None
    s=str(v).strip().strip("()")
    m=re.match(r"^\s*(-?[\d.]+)\s*[/:]\s*(-?[\d.]+)\s*$",s)
    return (float(m.group(1)),float(m.group(2))) if m else None
def num(v):
    if v is None or v=="" : return None
    try: return float(str(v).replace("cc","").replace(",",""))
    except: return None
diffs=collections.defaultdict(list); checked=collections.Counter()
def chk(field,name,a,b,tol=0.51):
    checked[field]+=1
    if a is None and b is None: return
    if a is None or b is None or abs(a-b)>tol: diffs[field].append((name,a,b))
SPEC=[("Weight (w/ Stage 3)","s3w"),("Wind Drag","wDrag"),("Grip Modifier (F)","gF"),("Grip Modifier (R)","gR"),("Wheelbase","wb"),
 ("Tyre Height (F)","thF"),("Tyre Height (R)","thR"),("Flywheel Inertia","fwI"),("Engine Brake","eBrk"),("Drive Inertia (F)","dIF"),("Drive Inertia (R)","dIR"),
 ("Prop. Inertia (F)","pIF"),("Prop. Inertia (R)","pIR"),("Yaw Radius","yaw"),("NOS Power","nos"),("Length","lng"),("Peak Torque (kgf.m)","pTq"),
 ("Chassis Tread (F)","trF"),("Chassis Tread (R)","trR"),("Chassis Width","wid")]
for n in g:
    c=cars[n]; s=specs.get(n); e=eng.get(n,{}); u=sus.get(n,{}); G=g[n]
    if s:
        for col,k in SPEC: chk("spec:"+k,n,num(c.get(col)),s.get(k),0.06)
        wd=pair(c.get("Weight Distribution")); chk("spec:wdF",n,wd and wd[0],s.get("wdF"))
        df=pair(c.get("Max DF (F / R)")); chk("spec:dfF",n,df and df[0],s.get("df",[None])[0]); chk("spec:dfR",n,df and df[1],(s.get("df")+[None,None])[1] if s.get("df") else None)
        for i,col in enumerate(["LSD Initial (F / R)","LSD Accel (F / R)","LSD Decel (F / R)"]):
            p=pair(c.get(col)); l=s.get("lsd") or [None]*6
            chk("spec:lsd"+str(i)+"F",n,p and p[0],l[2*i]); chk("spec:lsd"+str(i)+"R",n,p and p[1],l[2*i+1])
        for col,k in (("Tyre Size (F)","tsF"),("Tyre Size (R)","tsR"),("Aspiration","asp")):
            checked["spec:"+k]+=1
            if str(c.get(col)).strip()!=str(s.get(k)).strip(): diffs["spec:"+k].append((n,c.get(col),s.get(k)))
        checked["spec:canW"]+=1
        if (str(c.get("Can Equip Wing?")).lower()=="yes")!=(s.get("canW")==1): diffs["spec:canW"].append((n,c.get("Can Equip Wing?"),s.get("canW")))
        checked["spec:canSC"]+=1
        if (str(c.get("Can Equip S/C?")).lower()=="yes")!=(s.get("canSC")==1): diffs["spec:canSC"].append((n,c.get("Can Equip S/C?"),s.get("canSC")))
        # engine sheet
        chk("spec:redl",n,num(e.get("Redline")),s.get("redl")); chk("spec:revL",n,num(e.get("Rev Limiter")),s.get("revL"))
        checked["spec:disp"]+=1
        if str(e.get("Displacement")).strip()!=str(s.get("disp")).strip(): diffs["spec:disp"].append((n,e.get("Displacement"),s.get("disp")))
        # torque points
        pts=[]
        for k,v in e.items():
            if k.startswith("torquePoint") and v:
                m=re.match(r"\s*([\d.]+)\s*\(@\s*([\d.]+)\)",str(v))
                if m: pts.append((float(m.group(2)),float(m.group(1))))
        tc=s.get("tc") or []
        if tc and isinstance(tc[0],list): spts=[(a,b) for a,b in tc]
        elif tc: spts=[(tc[0]+i*tc[1],v) for i,v in enumerate(tc[2:])]
        else: spts=[]
        checked["spec:tc"]+=1
        sd=dict(spts); bad=[(r,v,sd.get(r)) for r,v in pts if sd.get(r) is None or abs(sd[r]-v)>0.6]
        if bad or len(pts)!=len(spts): diffs["spec:tc"].append((n,len(pts),len(spts),bad[:3]))
        # suspension
        sp=s.get("susp") or []
        cols=["Stock Springs Min","Stock Springs Max","Custom Springs Min","Custom Springs Max","Stock Height Min","Stock Height Max","Custom Height Min","Custom Height Max"]
        for i,col in enumerate(cols):
            p=pair(u.get(col))
            if len(sp)==16: chk("spec:susp:"+col+" F",n,p and p[0],sp[2*i],0.06); chk("spec:susp:"+col+" R",n,p and p[1],sp[2*i+1],0.06)
    # game json vs sheet
    chk("game:hp vs ps",n,num(c.get("Peak Power (ps)")),G.get("hp"))
    chk("game:widthMm",n,num(c.get("Chassis Width")) or 0,G.get("widthMm")); chk("game:trackFMm",n,num(c.get("Chassis Tread (F)")) or 0,G.get("trackFMm")); chk("game:trackRMm",n,num(c.get("Chassis Tread (R)")) or 0,G.get("trackRMm"))
    chk("game:modelYear",n,num(c.get("Year")),G.get("modelYear")); chk("game:redline",n,num(e.get("Redline")),G.get("redline"))
    chk("game:dispCc",n,num(e.get("Displacement")),G.get("dispCc"))
    pt=num(c.get("Peak Torque (kgf.m)")); chk("game:peakTorqueNm",n,pt and round(pt*9.80665),G.get("peakTorqueNm"),1.5)
    checked["game:drv"]+=1
    if str(c.get("Drivetrain")).strip()!=G.get("drv"): diffs["game:drv"].append((n,c.get("Drivetrain"),G.get("drv")))
    checked["game:asp"]+=1
    if str(c.get("Aspiration")).strip()!=G.get("asp"): diffs["game:asp"].append((n,c.get("Aspiration"),G.get("asp")))
    checked["game:eType"]+=1
    if str(e.get("Engine Type")).strip()!=G.get("eType"): diffs["game:eType"].append((n,e.get("Engine Type"),G.get("eType")))
    chk("game:kg vs s3w",n,num(c.get("Weight (w/ Stage 3)")),G.get("minKg"))
print("game cars",len(g),"with RG2 spec",sum(1 for n in g if n in specs))
for f in sorted(checked):
    d=diffs.get(f,[])
    print(f"{f:40s} checked {checked[f]:4d}  differ {len(d):4d}  " + ("; ".join(str(x)[:110] for x in d[:2]) if d else ""))

out = {"cars": cars, "engine": eng, "suspension": sus}
with io.open(os.path.join(HERE, "gt4_specs.json"), "w", encoding="utf-8") as f:
    json.dump(out, f, ensure_ascii=False, indent=0, default=str)
print("wrote", os.path.join(HERE, "gt4_specs.json"))
